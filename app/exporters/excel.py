"""프로그램 목록 Excel 산출물 생성 (openpyxl).

정해진 컬럼 순서(models.COLUMNS)대로 헤더/데이터를 채운다.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import COLUMNS

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
CELL_FONT = Font(size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def build_program_list(programs) -> bytes:
    """프로그램 목록을 xlsx 바이트로 반환."""
    wb = Workbook()
    ws = wb.active
    ws.title = "프로그램목록"

    labels = ["No"] + [label for _, label, _ in COLUMNS]

    # 헤더
    for col_idx, label in enumerate(labels, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    # 데이터
    for row_idx, program in enumerate(programs, start=2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1).alignment = CENTER
        data = program.as_dict()
        for col_idx, (name, _, typ) in enumerate(COLUMNS, start=2):
            value = data.get(name)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = CELL_FONT
            cell.border = BORDER
            cell.alignment = LEFT if typ in ("text", "textarea") else CENTER
        ws.cell(row=row_idx, column=1).border = BORDER

    # 열 너비 대략 조정
    widths = [5] + [16] * len(COLUMNS)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(labels))}1"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
