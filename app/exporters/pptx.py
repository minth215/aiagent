"""업무흐름도 PPT 산출물 생성 (python-pptx).

두 종류의 흐름을 그린다:
  1) 업무 프로세스 흐름 : 업무분류별로 프로그램레벨1을 순서대로 나열 (예: 신청 → 승인 → 정산)
  2) 호출 흐름         : 프로그램유형을 계층으로 매핑 (화면 → Controller → Service → Mapper → DB)

프로그램유형 값은 사용자가 자유 입력하므로 키워드로 계층을 추정한다.
매칭되지 않으면 '기타' 계층으로 분류한다.
"""
import io
from collections import OrderedDict, defaultdict

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# 호출 흐름 계층 정의: (라벨, 색상, 매칭 키워드들)
LAYERS = [
    ("화면(JSP)", "2E75B6", ["jsp", "화면", "screen", "view", "ui", "html"]),
    ("Controller", "5B9BD5", ["controller", "컨트롤러", "ctrl", "servlet"]),
    ("Service", "70AD47", ["service", "서비스", "svc", "biz", "business", "로직"]),
    ("Mapper/DAO", "ED7D31", ["mapper", "dao", "sql", "mybatis", "쿼리", "query"]),
    ("배치/기타", "A5A5A5", ["batch", "배치", "sh", "shell", "job", "scheduler"]),
    ("DB/테이블", "7030A0", ["table", "테이블", "db", "database", "entity"]),
]
ETC_LAYER = "기타"


def _layer_of(program_type: str) -> str:
    t = (program_type or "").lower()
    for label, _, keywords in LAYERS:
        if any(k in t for k in keywords):
            return label
    return ETC_LAYER


def _color(hex_str: str) -> RGBColor:
    return RGBColor.from_string(hex_str)


def _add_box(slide, left, top, width, height, text, fill="4472C4", font_size=11, font_color="FFFFFF"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(fill)
    shape.line.color.rgb = _color("FFFFFF")
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Pt(4)
    tf.margin_right = Pt(4)
    tf.margin_top = Pt(2)
    tf.margin_bottom = Pt(2)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = _color(font_color)
    run.font.bold = True
    return shape


def _add_arrow(slide, left, top, width, height, direction="right"):
    shape_type = MSO_SHAPE.RIGHT_ARROW if direction == "right" else MSO_SHAPE.DOWN_ARROW
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color("BFBFBF")
    shape.line.fill.background()
    return shape


def _add_title(slide, text, subtitle=None):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(26)
    run.font.bold = True
    run.font.color.rgb = _color("1F3864")
    if subtitle:
        p2 = tf.add_paragraph()
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(13)
        r2.font.color.rgb = _color("808080")


def _blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # 빈 레이아웃


def _priority_key(value):
    """우선순위가 숫자면 숫자로, 아니면 문자열로 정렬."""
    if value is None:
        return (1, "")
    s = str(value).strip()
    try:
        return (0, f"{int(s):09d}")
    except ValueError:
        return (1, s)


def _process_flow_slide(prs, biz_category, programs):
    """업무 프로세스 흐름: 프로그램레벨1을 순서대로 나열."""
    slide = _blank_slide(prs)
    _add_title(slide, f"업무 프로세스 흐름 — {biz_category}", "프로그램레벨1 기준 진행 순서")

    # 레벨1별 그룹, 각 그룹의 대표 우선순위(최소값)로 정렬
    groups = defaultdict(list)
    for p in programs:
        groups[p.program_level1 or "(미분류)"].append(p)

    def group_priority(items):
        return min((_priority_key(i.priority) for i in items), default=(1, ""))

    ordered = sorted(groups.items(), key=lambda kv: group_priority(kv[1]))

    n = len(ordered)
    if n == 0:
        return
    max_per_row = 4
    box_w = Inches(2.6)
    box_h = Inches(0.9)
    gap = Inches(0.5)
    start_x = Inches(0.6)
    start_y = Inches(1.7)
    row_h = Inches(2.0)

    for idx, (level1, items) in enumerate(ordered):
        row = idx // max_per_row
        col = idx % max_per_row
        x = start_x + col * (box_w + gap)
        y = start_y + row * row_h

        _add_box(slide, x, y, box_w, box_h, level1, fill="1F4E78", font_size=13)

        # 하위 프로그램 목록 (최대 6개)
        names = [i.program_name or i.file_name or "-" for i in items]
        listing = "\n".join(f"· {name}" for name in names[:6])
        if len(names) > 6:
            listing += f"\n  … 외 {len(names) - 6}건"
        detail = slide.shapes.add_textbox(x, y + box_h + Inches(0.05), box_w, Inches(0.95))
        tf = detail.text_frame
        tf.word_wrap = True
        first = True
        for line in listing.split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            r = p.add_run()
            r.text = line
            r.font.size = Pt(9)
            r.font.color.rgb = _color("404040")

        # 화살표 (같은 행 내 다음 박스로)
        if col < max_per_row - 1 and idx < n - 1:
            _add_arrow(slide, x + box_w + Inches(0.05), y + Inches(0.25), gap - Inches(0.1), Inches(0.4))


def _call_flow_slide(prs, biz_category, programs):
    """호출 흐름: 프로그램유형을 계층으로 매핑."""
    slide = _blank_slide(prs)
    _add_title(slide, f"호출 흐름 — {biz_category}", "프로그램유형 기준 계층 (화면 → Controller → Service → Mapper → DB)")

    layer_order = [label for label, _, _ in LAYERS] + [ETC_LAYER]
    layer_color = {label: color for label, color, _ in LAYERS}
    layer_color[ETC_LAYER] = "808080"

    buckets = OrderedDict((label, []) for label in layer_order)
    for p in programs:
        buckets[_layer_of(p.program_type)].append(p)

    active = [(label, items) for label, items in buckets.items() if items]
    if not active:
        return

    n = len(active)
    col_gap = Inches(0.3)
    total_w = Inches(12.5)
    col_w = Emu(int((total_w - col_gap * (n - 1)) / n))
    start_x = Inches(0.4)
    header_y = Inches(1.6)
    header_h = Inches(0.6)
    body_y = Inches(2.3)

    for idx, (label, items) in enumerate(active):
        x = start_x + idx * (col_w + col_gap)
        # 계층 헤더
        _add_box(slide, x, header_y, col_w, header_h, label, fill=layer_color[label], font_size=13)

        # 계층 내 프로그램 박스들
        item_h = Inches(0.55)
        item_gap = Inches(0.15)
        y = body_y
        for item in items[:7]:
            name = item.program_name or item.file_name or "-"
            _add_box(slide, x, y, col_w, item_h, name, fill="D9E1F2", font_size=9, font_color="1F3864")
            y = y + item_h + item_gap
        if len(items) > 7:
            extra = slide.shapes.add_textbox(x, y, col_w, Inches(0.3))
            r = extra.text_frame.paragraphs[0].add_run()
            r.text = f"… 외 {len(items) - 7}건"
            r.font.size = Pt(9)
            r.font.color.rgb = _color("808080")

        # 계층 간 화살표
        if idx < n - 1:
            _add_arrow(slide, x + col_w + Inches(0.02), header_y + Inches(0.15), col_gap - Inches(0.04), Inches(0.3))


def build_flow_diagram(programs) -> bytes:
    """업무흐름도(PPT)를 바이트로 반환."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 표지
    cover = _blank_slide(prs)
    _add_title(cover, "업무흐름도", f"총 {len(programs)}개 프로그램 · 업무 프로세스 흐름 + 호출 흐름")

    # 업무분류별로 슬라이드 생성
    by_category = defaultdict(list)
    for p in programs:
        by_category[p.biz_category or "(미분류)"].append(p)

    for biz_category in sorted(by_category.keys()):
        items = by_category[biz_category]
        _process_flow_slide(prs, biz_category, items)
        _call_flow_slide(prs, biz_category, items)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
